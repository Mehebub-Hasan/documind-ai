from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


SECTION_TARGET_CHARS = 2000


class DocxExtractionError(Exception):
    """Raised when a file cannot be read as a valid .docx document."""


def _iter_block_items(document: DocumentObject):
    """
    Yield Paragraph and Table objects in true document order.

    python-docx exposes document.paragraphs and document.tables as separate
    lists, which loses ordering. Iterating the body XML keeps the original
    sequence so source references stay accurate.
    """

    body = document.element.body

    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _table_to_text(table: Table) -> str:
    """
    Flatten a table into readable text: cells joined with " | " per row,
    rows joined with newlines. Nested tables are appended to their cell.
    """

    rows = []

    for row in table.rows:
        cells = []

        for cell in row.cells:
            cell_parts = [
                paragraph.text.strip()
                for paragraph in cell.paragraphs
                if paragraph.text.strip()
            ]

            for nested_table in cell.tables:
                nested_text = _table_to_text(nested_table).strip()
                if nested_text:
                    cell_parts.append(nested_text)

            cells.append("\n".join(cell_parts))

        rows.append(" | ".join(cells))

    return "\n".join(rows)


def extract_blocks_from_docx(docx_path: str):
    """
    Extract paragraphs and tables from a .docx file in document order.

    Returns:
        list of dictionaries:
        [
            {"block": 1, "type": "paragraph", "text": "..."},
            {"block": 2, "type": "table", "text": "cell | cell"},
        ]

        "block" is the 1-based position in document order. Word documents
        have no fixed pages, so blocks are the source unit instead.
    """

    docx_file = Path(docx_path)

    if not docx_file.exists():
        raise FileNotFoundError(f"DOCX file not found: {docx_path}")

    try:
        document = Document(str(docx_path))
    except Exception as exc:
        raise DocxExtractionError(
            "The file could not be read as a valid .docx document. "
            "It may be corrupt or not a Word file."
        ) from exc

    blocks = []

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            block_type = "paragraph"
        else:
            text = _table_to_text(block).strip()
            block_type = "table"

        blocks.append(
            {
                "block": len(blocks) + 1,
                "type": block_type,
                "text": text
            }
        )

    return blocks


def group_blocks_into_sections(blocks, target_chars: int = SECTION_TARGET_CHARS):
    """
    Group consecutive non-empty blocks into sections of roughly target_chars
    characters, so chunk density stays comparable to PDF pages.

    A single block longer than target_chars becomes its own section.

    Returns:
        list of dictionaries:
        [
            {"section": 1, "block_start": 1, "block_end": 9, "text": "..."},
        ]
    """

    sections = []
    current_blocks = []
    current_texts = []
    current_length = 0

    def close_section():
        if current_blocks:
            sections.append(
                {
                    "section": len(sections) + 1,
                    "block_start": current_blocks[0],
                    "block_end": current_blocks[-1],
                    "text": "\n\n".join(current_texts).strip()
                }
            )

    for block in blocks:
        text = block["text"]
        if not text:
            continue

        if current_blocks and current_length + len(text) > target_chars:
            close_section()
            current_blocks = []
            current_texts = []
            current_length = 0

        current_blocks.append(block["block"])
        current_texts.append(text)
        current_length += len(text)

    close_section()

    return sections


def get_docx_summary(blocks):
    """
    Create a simple summary of extracted docx blocks.
    """

    total_blocks = len(blocks)
    total_paragraphs = sum(1 for block in blocks if block["type"] == "paragraph")
    total_tables = sum(1 for block in blocks if block["type"] == "table")
    total_characters = sum(len(block["text"]) for block in blocks)

    preview_text = ""

    for block in blocks:
        if block["text"].strip():
            preview_text = block["text"][:500]
            break

    return {
        "total_blocks": total_blocks,
        "total_paragraphs": total_paragraphs,
        "total_tables": total_tables,
        "total_characters": total_characters,
        "preview_text": preview_text
    }


def create_units_from_sections(sections):
    """
    Convert sections into generic chunker units with block-range metadata.

    Word files have no reliable page numbers, so chunks carry the document
    block range they came from instead.
    """

    units = []

    for section in sections:
        start = section["block_start"]
        end = section["block_end"]

        label = f"block_{start}" if start == end else f"blocks_{start}-{end}"

        units.append(
            {
                "label": label,
                "text": section["text"],
                "meta": {"block_start": start, "block_end": end}
            }
        )

    return units
