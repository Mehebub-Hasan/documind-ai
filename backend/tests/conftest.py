import io
import os
import sys
import tempfile
import uuid
import zlib
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

# Isolated test environment. These must be set before `main` is imported,
# because database.py / document_store.py / main.py read them at import time.
TEST_ROOT = Path(tempfile.mkdtemp(prefix="documind_tests_"))
os.environ.setdefault("DATABASE_URL", f"sqlite:///{TEST_ROOT / 'documind_test.db'}")
os.environ.setdefault("JWT_SECRET_KEY", "documind-test-secret-that-is-long-enough-0123456789")
os.environ.setdefault("INDEX_STORAGE_DIR", str(TEST_ROOT / "indexes"))
os.environ.setdefault("UPLOAD_DIR", str(TEST_ROOT / "uploads"))
os.environ.setdefault("BCRYPT_ROUNDS", "4")  # fast hashes for tests

import numpy as np
import pytest
from fastapi.testclient import TestClient

import fitz
from docx import Document

import document_store
from auth import create_access_token
from database import SessionLocal, engine
from main import app, UPLOAD_DIR
import main as main_module
from models import ChatMessage, Document as DocumentModel, DocumentChunk, User
from vector_store import VectorStore

EMBEDDING_DIM = 8

# Captured before any test patches the class; tests that exercise the real
# embedding dispatch re-apply this via the real_embed_texts fixture.
ORIGINAL_EMBED_TEXTS = VectorStore._embed_texts


def _stub_embed_texts(self, texts):
    """
    Deterministic offline embeddings so tests never call the OpenAI API.
    Uses crc32 (stable across processes) rather than Python's hash().
    """

    vectors = []

    for text in texts:
        vector = [0.0] * EMBEDDING_DIM

        for token in text.lower().split():
            vector[zlib.crc32(token.encode("utf-8")) % EMBEDDING_DIM] += 1.0

        if not any(vector):
            vector[0] = 1.0

        vectors.append(vector)

    return np.array(vectors, dtype="float32")


@pytest.fixture(autouse=True)
def stub_embeddings(monkeypatch):
    monkeypatch.setattr(VectorStore, "_embed_texts", _stub_embed_texts)
    monkeypatch.setattr(
        VectorStore, "embedding_dimension", lambda self: EMBEDDING_DIM
    )


@pytest.fixture(autouse=True)
def disable_llm_calls(monkeypatch):
    """
    Keep the suite hermetic: a developer's backend/.env may contain real API
    keys, and uploads generate executive summaries through the LLM. Tests
    that exercise the LLM paths patch in their own fake clients.
    """

    monkeypatch.setattr("summarizer.get_deepseek_client", lambda: None)
    monkeypatch.setattr("chatbot.get_deepseek_client", lambda: None)


@pytest.fixture(autouse=True)
def relax_similarity_gate(monkeypatch):
    """
    Disable the FAISS similarity gate by default: the deterministic stub
    embeddings produce far lower cosine scores than real embedding models,
    so most API tests would be dropped by the guardrail. Tests that verify
    the gate itself re-enable it with monkeypatch on main.MIN_SIMILARITY_SCORE.
    """

    monkeypatch.setattr(main_module, "MIN_SIMILARITY_SCORE", -1.0)


@pytest.fixture(autouse=True)
def clean_state():
    """
    Reset in-memory caches, then wipe the database tables and generated
    files after every test so tests stay fully independent.
    """

    document_store.clear_cache()
    yield
    document_store.clear_cache()

    with engine.begin() as connection:
        for table in (
            ChatMessage.__table__,
            DocumentChunk.__table__,
            DocumentModel.__table__,
            User.__table__,
        ):
            connection.execute(table.delete())

    for folder in (UPLOAD_DIR, document_store.INDEX_STORAGE_DIR):
        for item in folder.glob("*"):
            if item.is_file():
                item.unlink()


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture(autouse=True)
def disable_email_delivery(monkeypatch):
    """
    Tests never deliver email: send_verification_email is stubbed out even
    if a developer's .env carries a real RESEND_API_KEY. Codes land in the
    database and tests assert on those.

    Patched in BOTH namespaces: utils.email_service (the definition) and
    main (which holds its own imported reference and hands it to
    BackgroundTasks.add_task at request time).
    """

    stub = lambda email, code: False  # noqa: E731
    monkeypatch.setattr(
        "utils.email_service.send_verification_email", stub
    )
    monkeypatch.setattr("main.send_verification_email", stub)


def make_user(client, email=None, password="Password123!", name="Test User"):
    """
    Register a fresh account, mark it verified directly in the database
    (tests skip the email round-trip), and mint a token. Returns
    credentials, token, auth headers, and the public profile.
    """

    email = email or f"user-{uuid.uuid4().hex[:12]}@example.com"

    response = client.post(
        "/auth/register",
        json={"email": email, "password": password, "name": name},
    )

    assert response.status_code == 200, response.text

    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).first()
        assert user is not None

        user.is_verified = True
        db.commit()
        db.refresh(user)

        token = create_access_token(user)
        profile = {
            "id": user.id,
            "email": user.email,
            "name": user.name,
            "is_verified": bool(user.is_verified),
        }
    finally:
        db.close()

    return {
        "email": email,
        "password": password,
        "name": name,
        "token": token,
        "user": profile,
        "headers": {"Authorization": f"Bearer {token}"},
    }


def build_pdf_bytes(pages_text):
    """
    Create a small in-memory PDF with one page per text string.
    """

    with fitz.open() as pdf:
        for text in pages_text:
            page = pdf.new_page()
            page.insert_text((72, 72), text)
        return pdf.tobytes()


def build_docx_bytes(items=()):
    """
    Create a small in-memory .docx from ordered items.

    Each item is ("p", "paragraph text") or ("t", [["c1", "c2"], ...]).
    """

    document = Document()

    for kind, payload in items:
        if kind == "p":
            document.add_paragraph(payload)
        elif kind == "t":
            rows = payload
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for col_index, cell_text in enumerate(row):
                    table.rows[row_index].cells[col_index].text = cell_text

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def pdf_bytes():
    return build_pdf_bytes(
        [
            "DocuMind test document. Alpha section with unique token apple.",
            "Beta section with unique token banana.",
        ]
    )


@pytest.fixture
def docx_bytes():
    return build_docx_bytes(
        [
            ("p", "Intro paragraph with unique token apple."),
            ("p", "Second paragraph explaining the quarterly results in detail."),
            ("p", "Closing paragraph with unique token banana."),
            ("t", [["Item", "Value"], ["Revenue", "1200"]]),
        ]
    )
