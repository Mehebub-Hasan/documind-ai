import pytest
from docx import Document

from docx_processor import (
    DocxExtractionError,
    create_units_from_sections,
    extract_blocks_from_docx,
    get_docx_summary,
    group_blocks_into_sections,
)
from conftest import build_docx_bytes


def make_blocks(texts):
    return [
        {"block": index + 1, "type": "paragraph", "text": text}
        for index, text in enumerate(texts)
    ]


class TestExtractBlocks:

    def test_paragraph_and_table_document_order(self, tmp_path):
        path = tmp_path / "ordered.docx"
        path.write_bytes(
            build_docx_bytes(
                [
                    ("p", "First paragraph."),
                    ("t", [["h1", "h2"], ["a", "b"]]),
                    ("p", "After table."),
                ]
            )
        )

        blocks = extract_blocks_from_docx(str(path))

        assert [block["type"] for block in blocks] == [
            "paragraph",
            "table",
            "paragraph",
        ]
        assert [block["block"] for block in blocks] == [1, 2, 3]
        assert blocks[0]["text"] == "First paragraph."
        assert "h1 | h2" in blocks[1]["text"]
        assert "a | b" in blocks[1]["text"]
        assert blocks[2]["text"] == "After table."

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            extract_blocks_from_docx(str(tmp_path / "missing.docx"))

    def test_invalid_file_raises_docx_error(self, tmp_path):
        path = tmp_path / "broken.docx"
        path.write_bytes(b"this is not a zip archive")

        with pytest.raises(DocxExtractionError):
            extract_blocks_from_docx(str(path))

    def test_empty_document_yields_no_sections(self, tmp_path):
        path = tmp_path / "empty.docx"
        path.write_bytes(build_docx_bytes())

        blocks = extract_blocks_from_docx(str(path))
        sections = group_blocks_into_sections(blocks)

        assert sections == []

    def test_nested_table_extracted(self, tmp_path):
        document = Document()
        document.add_paragraph("With nested table")

        table = document.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        cell.text = "outer cell"

        nested = cell.add_table(rows=1, cols=1)
        nested.rows[0].cells[0].text = "inner cell"

        path = tmp_path / "nested.docx"
        document.save(str(path))

        blocks = extract_blocks_from_docx(str(path))

        assert [block["type"] for block in blocks] == ["paragraph", "table"]
        assert "outer cell" in blocks[1]["text"]
        assert "inner cell" in blocks[1]["text"]


class TestGroupBlocksIntoSections:

    def test_groups_consecutive_blocks(self):
        blocks = make_blocks(["x" * 100 for _ in range(5)])

        sections = group_blocks_into_sections(blocks, target_chars=250)

        assert [
            (section["block_start"], section["block_end"]) for section in sections
        ] == [(1, 2), (3, 4), (5, 5)]
        assert all(section["text"].startswith("x") for section in sections)

    def test_oversized_block_becomes_own_section(self):
        blocks = make_blocks(["y" * 5000, "short text"])

        sections = group_blocks_into_sections(blocks, target_chars=250)

        assert [
            (section["block_start"], section["block_end"]) for section in sections
        ] == [(1, 1), (2, 2)]

    def test_empty_blocks_are_skipped(self):
        blocks = [
            {"block": 1, "type": "paragraph", "text": ""},
            {"block": 2, "type": "paragraph", "text": "real text"},
        ]

        sections = group_blocks_into_sections(blocks)

        assert len(sections) == 1
        assert sections[0]["block_start"] == 2
        assert sections[0]["block_end"] == 2


class TestGetDocxSummary:

    def test_summary_counts(self):
        blocks = [
            {"block": 1, "type": "paragraph", "text": "alpha"},
            {"block": 2, "type": "table", "text": "a | b"},
        ]

        summary = get_docx_summary(blocks)

        assert summary == {
            "total_blocks": 2,
            "total_paragraphs": 1,
            "total_tables": 1,
            "total_characters": 10,
            "preview_text": "alpha",
        }


class TestCreateUnitsFromSections:

    def test_labels_and_metadata(self):
        sections = [
            {"section": 1, "block_start": 4, "block_end": 9, "text": "t"},
            {"section": 2, "block_start": 12, "block_end": 12, "text": "u"},
        ]

        units = create_units_from_sections(sections)

        assert units[0]["label"] == "blocks_4-9"
        assert units[0]["meta"] == {"block_start": 4, "block_end": 9}
        assert units[1]["label"] == "block_12"
        assert units[1]["meta"] == {"block_start": 12, "block_end": 12}
